# -*- coding: utf-8 -*-
import os
import time
import pickle
from typing import NamedTuple, List, Optional

import jax
import jax.numpy as jnp
import numpy as np
import optax
import distrax
from flax import linen as nn
from flax.training.train_state import TrainState

# 设置 Matplotlib 后端为 Agg，防止在无显示器的服务器上报错
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# ==========================================
# 0. 平滑工具函数 & 环境导入
# ==========================================
try:
    from scipy.ndimage import gaussian_filter1d
except ImportError:
    gaussian_filter1d = None

def smooth_curve(
    x: np.ndarray,
    method: str = "gaussian",
    sigma: float = 5.0,
    alpha: float = 0.1,
    window: int = 10,
    passes: int = 1,
) -> np.ndarray:
    """
    对奖励曲线进行平滑处理 (仅用于绘图，不影响训练数据)
    支持: 'gaussian', 'ema' (指数移动平均), 'uniform' (滑动窗口)
    """
    x = np.asarray(x, dtype=np.float32)
    if x.size <= 2:
        return x
    method = (method or "").lower()
    passes = max(int(passes), 1)
    y = x.copy()
    
    for _ in range(passes):
        method_local = method
        # 1. 高斯滤波 (首选，最平滑且自然)
        if method_local == "gaussian":
            if gaussian_filter1d is not None and sigma > 0:
                y = gaussian_filter1d(y, float(sigma), mode="nearest")
                continue
            method_local = "ema" # 降级处理
            
        # 2. 指数移动平均 (EMA)
        if method_local in ("ema", "exponential"):
            a = float(alpha)
            a = min(max(a, 0.0), 0.999)
            y_next = np.empty_like(y)
            y_next[0] = y[0]
            for i in range(1, y.size):
                y_next[i] = a * y_next[i - 1] + (1.0 - a) * y[i]
            y = y_next
            continue
            
        # 3. 滑动窗口平均
        if method_local in ("uniform", "moving_average", "ma") and window > 1:
            w = int(window)
            kernel = np.ones((w,), dtype=np.float32) / float(w)
            pad = w // 2
            yp = np.pad(y, (pad, w - 1 - pad), mode="edge")
            y = np.convolve(yp, kernel, mode="valid")
            continue
            
        return y
    return y

try:
    from rl_env.path_env import RlGame
except ImportError:
    try:
        from path_env import RlGame
    except ImportError:
        print("Error: 找不到 'path_env.py'。请确保该文件在当前目录下或在 PYTHONPATH 中。")
        exit(1)

# ==========================================
# 1. 配置 (Config)
# ==========================================
class Config:
    # 模式开关
    TEST_MODE = False    # True: 进行测试并输出指标; False: 进行训练
    TEST_MODEL_PATH = "checkpoints/dsact_jax_final.pkl" # 模型路径

    # Environment
    N_AGENT = 1
    M_ENEMY = 1
    RENDER = False      
    TEST_RENDER = False  
    
    # Training
    SEED = 42
    EP_MAX = 50000      
    EP_LEN = 1000       
    BATCH_SIZE = 128
    BUFFER_SIZE = 100000
    GAMMA = 0.99
    
    # Optimization
    ACTOR_LR = 3e-4
    CRITIC_LR = 3e-4
    ALPHA_LR = 3e-4
    TAU = 0.01          
    
    # DSAC-T Hyperparameters
    TAU_B = 0.01        
    STD_BIAS = 0.1
    HUBER_DELTA = 50.0
    TD_BOUND_SCALE = 3.0
    
    # Updates
    START_STEPS = 2000
    UPDATE_EVERY = 1
    DELAY_UPDATE = 2
    
    # Saving
    SAVE_DIR = "./checkpoints"
    SAVE_EVERY = 25000 

    # [新增] 导出每个 Episode 的奖励到 Excel
    EXPORT_REWARDS_EXCEL = True
    REWARDS_EXCEL_NAME = "rewards_per_episode.xlsx"
    
    # Testing Metrics
    TEST_EPISODES = 100 

    # [新增] 绘图平滑参数
    SMOOTH_METHOD: str = "ema"         # ema | gaussian | uniform | none
    SMOOTH_SIGMA: float = 5.0          # 高斯滤波的 sigma
    SMOOTH_ALPHA: float = 0.95         # EMA 平滑系数 (越大越平滑)
    SMOOTH_WINDOW: int = 20            # 滑动窗口大小
    SMOOTH_PASSES: int = 2             # 平滑重复次数

# ==========================================
# 2. 工具函数 & Buffer
# ==========================================
def unwrap_step(step_out):
    if isinstance(step_out, (list, tuple)):
        obs = step_out[0]
        r = step_out[1]
        done = step_out[2]
        win = step_out[3] if len(step_out) > 3 else False
        team_counter = step_out[4] if len(step_out) > 4 else 0
        return obs, r, done, win, team_counter
    return step_out

def save_model(agents, filepath):
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    params_list = [jax.device_get(agent.actor.params) for agent in agents]
    with open(filepath, 'wb') as f:
        pickle.dump(params_list, f)
    print(f"模型已保存: {filepath}")

def load_model(agents, filepath):
    print(f"正在加载模型: {filepath} ...")
    with open(filepath, 'rb') as f:
        params_list = pickle.load(f)
    
    train_num = len(params_list)
    
    new_agents = []
    for i, agent in enumerate(agents):
        if i < train_num:
            target_params = params_list[i]
        else:
            print(f"  - Agent {i} (Extra Follower) 复用 Agent {train_num - 1} 的参数")
            target_params = params_list[train_num - 1]
            
        new_actor = agent.actor.replace(params=target_params)
        new_agents.append(agent._replace(actor=new_actor))
        
    print("模型加载完成！")
    return new_agents


def export_rewards_to_excel(save_path: str,
                           total_rewards: List[float],
                           leader_rewards: List[float],
                           follower_rewards: List[float]) -> None:
    """把每个 episode 的奖励导出为 .xlsx (不影响训练流程)。"""
    os.makedirs(os.path.dirname(save_path), exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "rewards"

    headers = ["Episode", "TotalReward", "LeaderReward", "FollowerReward"]
    ws.append(headers)

    n = min(len(total_rewards), len(leader_rewards), len(follower_rewards))
    for ep_idx in range(n):
        tot = float(total_rewards[ep_idx])
        lead = float(leader_rewards[ep_idx])
        foll = follower_rewards[ep_idx]
        # follower 可能是 nan
        foll_val = None
        try:
            if foll is not None and not (isinstance(foll, float) and np.isnan(foll)):
                foll_val = float(foll)
        except Exception:
            pass
        ws.append([ep_idx + 1, tot, lead, foll_val])

    # 简单设置一下列宽，方便查看
    for col in range(1, len(headers) + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for cell in ws[col_letter]:
            v = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(v))
        ws.column_dimensions[col_letter].width = min(max(12, max_len + 2), 40)

    wb.save(save_path)
    print(f"✓ 已导出奖励到 Excel: {save_path}")

class ReplayBuffer:
    def __init__(self, capacity, obs_dim, act_dim, num_agents):
        self.capacity = capacity
        self.ptr = 0
        self.size = 0
        self.s = np.zeros((capacity, num_agents * obs_dim), dtype=np.float32)
        self.a = np.zeros((capacity, num_agents * act_dim), dtype=np.float32)
        self.r = np.zeros((capacity, num_agents), dtype=np.float32)
        self.d = np.zeros((capacity, num_agents), dtype=np.float32)
        self.s_next = np.zeros((capacity, num_agents * obs_dim), dtype=np.float32)

    def add(self, s, a, r, d, s_next):
        idx = self.ptr
        self.s[idx] = s.reshape(-1)
        self.a[idx] = a.reshape(-1)
        self.r[idx] = r.reshape(-1)
        self.d[idx] = d.reshape(-1)
        self.s_next[idx] = s_next.reshape(-1)
        self.ptr = (self.ptr + 1) % self.capacity
        self.size = min(self.size + 1, self.capacity)

    def sample(self, batch_size):
        idx = np.random.randint(0, self.size, size=batch_size)
        return (
            jnp.array(self.s[idx]),
            jnp.array(self.a[idx]),
            jnp.array(self.r[idx]),
            jnp.array(self.d[idx]),
            jnp.array(self.s_next[idx])
        )

# ==========================================
# 3. 网络定义 (Flax)
# ==========================================
class Actor(nn.Module):
    act_dim: int
    max_action: float = 1.0

    @nn.compact
    def __call__(self, x):
        x = nn.Dense(256)(x)
        x = nn.relu(x)
        x = nn.Dense(256)(x)
        x = nn.relu(x)
        mean = nn.Dense(self.act_dim)(x)
        log_std = nn.Dense(self.act_dim)(x)
        log_std = jnp.clip(log_std, -20, 2)
        
        base_dist = distrax.MultivariateNormalDiag(loc=mean, scale_diag=jnp.exp(log_std))
        bijector = distrax.Block(distrax.Tanh(), ndims=1)
        dist = distrax.Transformed(distribution=base_dist, bijector=bijector)
        return dist

class Critic(nn.Module):
    @nn.compact
    def __call__(self, s_global, a_all):
        x = jnp.concatenate([s_global, a_all], axis=-1)
        # Q1
        q1 = nn.Dense(512)(x)
        q1 = nn.relu(q1)
        q1 = nn.Dense(512)(q1)
        q1 = nn.relu(q1)
        q1_mean = nn.Dense(1)(q1)
        q1_log_std = nn.Dense(1)(q1)
        # Q2
        q2 = nn.Dense(512)(x)
        q2 = nn.relu(q2)
        q2 = nn.Dense(512)(q2)
        q2 = nn.relu(q2)
        q2_mean = nn.Dense(1)(q2)
        q2_log_std = nn.Dense(1)(q2)

        q1_log_std = jnp.clip(q1_log_std, -5.0, 2.0)
        q2_log_std = jnp.clip(q2_log_std, -5.0, 2.0)
        q1_std = nn.softplus(q1_log_std) + 1e-4
        q2_std = nn.softplus(q2_log_std) + 1e-4
        return q1_mean, q1_std, q2_mean, q2_std

# ==========================================
# 4. TrainStates 管理
# ==========================================
class AgentState(NamedTuple):
    actor: TrainState
    critic: TrainState
    target_critic: TrainState
    log_alpha: TrainState
    running_mean_std1: jnp.ndarray
    running_mean_std2: jnp.ndarray
    rng: jnp.ndarray

def create_agent_state(rng, obs_dim, act_dim, cfg: Config, num_agents):
    rng, actor_key, critic_key = jax.random.split(rng, 3)
    
    actor_def = Actor(act_dim)
    s_local_mock = jnp.ones((1, obs_dim))
    actor_params = actor_def.init(actor_key, s_local_mock)
    actor_state = TrainState.create(apply_fn=actor_def.apply, params=actor_params, tx=optax.adam(cfg.ACTOR_LR))
    
    global_dim = obs_dim * num_agents + act_dim * num_agents
    critic_def = Critic()
    s_global_mock = jnp.ones((1, obs_dim * num_agents))
    a_global_mock = jnp.ones((1, act_dim * num_agents))
    critic_params = critic_def.init(critic_key, s_global_mock, a_global_mock)
    
    critic_state = TrainState.create(apply_fn=critic_def.apply, params=critic_params, tx=optax.adam(cfg.CRITIC_LR))
    target_critic_state = TrainState.create(apply_fn=critic_def.apply, params=critic_params, tx=optax.set_to_zero())
    
    log_alpha = jnp.array([0.0], dtype=jnp.float32)
    alpha_state = TrainState.create(apply_fn=None, params=log_alpha, tx=optax.adam(cfg.ALPHA_LR))
    
    return AgentState(
        actor=actor_state, critic=critic_state, target_critic=target_critic_state,
        log_alpha=alpha_state, running_mean_std1=jnp.array(1.0), running_mean_std2=jnp.array(1.0), rng=rng
    )

# ==========================================
# 5. 核心算法逻辑 (JIT Compiled)
# ==========================================
@jax.jit
def huber_loss(pred, target, delta):
    diff = pred - target
    abs_diff = jnp.abs(diff)
    quadratic = jnp.minimum(abs_diff, delta)
    linear = abs_diff - quadratic
    return 0.5 * quadratic ** 2 + delta * linear

@jax.jit
def update_step(agent_states, batch, cfg_vals, global_step):
    b_s, b_a, b_r, b_d, b_s_next = batch
    num_agents = len(agent_states)
    new_agent_states = []
    
    obs_dim = b_s.shape[1] // num_agents
    act_dim = b_a.shape[1] // num_agents
    def get_slice(data, idx, dim): return data[:, idx*dim : (idx+1)*dim]

    next_actions_list = []
    next_log_probs_list = []
    current_rngs = []
    
    # 1. Target Policy
    for i in range(num_agents):
        state = agent_states[i]
        rng_now, key = jax.random.split(state.rng)
        current_rngs.append(rng_now)
        s_next_i = get_slice(b_s_next, i, obs_dim)
        dist = state.actor.apply_fn(state.actor.params, s_next_i)
        a_next, log_pi_next = dist.sample_and_log_prob(seed=key)
        next_actions_list.append(a_next)
        next_log_probs_list.append(log_pi_next)

    b_a_next_all = jnp.concatenate(next_actions_list, axis=-1)

    # 2. Update
    for i in range(num_agents):
        st = agent_states[i]
        s_i = get_slice(b_s, i, obs_dim)
        r_i = b_r[:, i:i+1]
        d_i = b_d[:, i:i+1]
        alpha = jnp.exp(st.log_alpha.params[0])
        
        def critic_loss_fn(c_params, run_std1, run_std2):
            q1_m, q1_s, q2_m, q2_s = st.critic.apply_fn(c_params, b_s, b_a)
            q1n_m, q1n_s, q2n_m, q2n_s = st.target_critic.apply_fn(st.target_critic.params, b_s_next, b_a_next_all)
            q_next_mean = jnp.minimum(q1n_m, q2n_m)
            target_q = r_i + cfg_vals['GAMMA'] * (1 - d_i) * (q_next_mean - alpha * next_log_probs_list[i])
            
            td_bound1 = cfg_vals['TD_BOUND_SCALE'] * run_std1
            target_q_bound1 = q1_m + jnp.clip(target_q - q1_m, -td_bound1, td_bound1)
            td_bound2 = cfg_vals['TD_BOUND_SCALE'] * run_std2
            target_q_bound2 = q2_m + jnp.clip(target_q - q2_m, -td_bound2, td_bound2)

            q1_s_det = jax.lax.stop_gradient(q1_s)
            q2_s_det = jax.lax.stop_gradient(q2_s)
            ratio1 = jnp.clip((run_std1**2)/(q1_s_det**2 + cfg_vals['STD_BIAS']), 0.1, 10.0)
            ratio2 = jnp.clip((run_std2**2)/(q2_s_det**2 + cfg_vals['STD_BIAS']), 0.1, 10.0)
            
            l1 = huber_loss(q1_m, jax.lax.stop_gradient(target_q), cfg_vals['HUBER_DELTA'])
            l1_var = q1_s_det**2 - huber_loss(jax.lax.stop_gradient(q1_m), jax.lax.stop_gradient(target_q_bound1), cfg_vals['HUBER_DELTA'])
            loss_q1 = ratio1 * (l1 + q1_s * l1_var / (q1_s_det + cfg_vals['STD_BIAS']))
            l2 = huber_loss(q2_m, jax.lax.stop_gradient(target_q), cfg_vals['HUBER_DELTA'])
            l2_var = q2_s_det**2 - huber_loss(jax.lax.stop_gradient(q2_m), jax.lax.stop_gradient(target_q_bound2), cfg_vals['HUBER_DELTA'])
            loss_q2 = ratio2 * (l2 + q2_s * l2_var / (q2_s_det + cfg_vals['STD_BIAS']))
            return jnp.mean(loss_q1 + loss_q2), (q1_s, q2_s)

        grad_c_fn = jax.value_and_grad(critic_loss_fn, has_aux=True)
        (c_loss, (curr_q1_s, curr_q2_s)), c_grads = grad_c_fn(st.critic.params, st.running_mean_std1, st.running_mean_std2)
        new_critic_state = st.critic.apply_gradients(grads=c_grads)
        
        new_run_std1 = (1 - cfg_vals['TAU_B']) * st.running_mean_std1 + cfg_vals['TAU_B'] * jnp.mean(curr_q1_s)
        new_run_std2 = (1 - cfg_vals['TAU_B']) * st.running_mean_std2 + cfg_vals['TAU_B'] * jnp.mean(curr_q2_s)

        def update_actor_alpha(actor_st, alpha_st, critic_p):
            target_entropy = -float(act_dim)
            def alpha_loss_fn(log_a_params):
                dist = actor_st.apply_fn(actor_st.params, s_i)
                _, log_pi_curr = dist.sample_and_log_prob(seed=current_rngs[i]) 
                return -jnp.mean(log_a_params * (log_pi_curr + target_entropy))

            alpha_l, alpha_grads = jax.value_and_grad(alpha_loss_fn)(alpha_st.params)
            new_alpha_st = alpha_st.apply_gradients(grads=alpha_grads)
            curr_alpha = jnp.exp(new_alpha_st.params[0])
            
            def actor_loss_fn(a_params):
                dist = actor_st.apply_fn(a_params, s_i)
                a_curr, log_pi_curr = dist.sample_and_log_prob(seed=current_rngs[i])
                b_a_replaced = b_a.at[:, i*act_dim:(i+1)*act_dim].set(a_curr)
                q1_pi, _, q2_pi, _ = st.critic.apply_fn(critic_p, b_s, b_a_replaced)
                return jnp.mean(curr_alpha * log_pi_curr - jnp.minimum(q1_pi, q2_pi))
            
            act_l, act_grads = jax.value_and_grad(actor_loss_fn)(actor_st.params)
            new_act_st = actor_st.apply_gradients(grads=act_grads)
            return new_act_st, new_alpha_st

        is_update = (global_step % cfg_vals['DELAY_UPDATE'] == 0)
        new_actor_state, new_alpha_state = jax.lax.cond(
            is_update, update_actor_alpha, lambda a,b,c: (a,b), st.actor, st.log_alpha, st.critic.params
        )

        new_target_params = optax.incremental_update(new_critic_state.params, st.target_critic.params, cfg_vals['TAU'])
        new_target_critic = st.target_critic.replace(params=new_target_params)

        new_agent_states.append(AgentState(
            actor=new_actor_state, critic=new_critic_state, target_critic=new_target_critic,
            log_alpha=new_alpha_state, running_mean_std1=new_run_std1, running_mean_std2=new_run_std2, rng=current_rngs[i]
        ))
        
    return new_agent_states

# ==========================================
# 6. 训练与测试逻辑
# ==========================================
def train(cfg: Config):
    print(f"初始化环境 (Train Mode): {cfg.N_AGENT} Agents ...")
    env = RlGame(n=cfg.N_AGENT, m=cfg.M_ENEMY, render=cfg.RENDER)
    if hasattr(env, 'unwrapped'): env = env.unwrapped
    
    obs = env.reset()
    obs = np.array(obs, dtype=np.float32)
    num_agents = cfg.N_AGENT + cfg.M_ENEMY
    obs_dim = obs.shape[-1]
    act_dim = env.action_space.shape[0]
    
    rng = jax.random.PRNGKey(cfg.SEED)
    agents = []
    for i in range(num_agents):
        rng, key = jax.random.split(rng)
        agents.append(create_agent_state(key, obs_dim, act_dim, cfg, num_agents))
    
    buffer = ReplayBuffer(cfg.BUFFER_SIZE, obs_dim, act_dim, num_agents)
    cfg_dict = {'GAMMA': cfg.GAMMA, 'TAU': cfg.TAU, 'TAU_B': cfg.TAU_B, 'STD_BIAS': cfg.STD_BIAS, 
                'HUBER_DELTA': cfg.HUBER_DELTA, 'TD_BOUND_SCALE': cfg.TD_BOUND_SCALE, 'DELAY_UPDATE': cfg.DELAY_UPDATE}
    
    global_step = 0
    # [新增] 分别记录奖励历史
    total_rewards_history = []
    leader_rewards_history = []
    follower_rewards_history = []
    
    print("开始训练...")
    for ep in range(cfg.EP_MAX):
        obs = env.reset()
        obs = np.array(obs, dtype=np.float32)
        ep_ret = np.zeros(num_agents)
        
        for t in range(cfg.EP_LEN):
            global_step += 1
            actions = []
            for i in range(num_agents):
                obs_tensor = jnp.expand_dims(jnp.array(obs[i]), 0)
                dist = agents[i].actor.apply_fn(agents[i].actor.params, obs_tensor)
                rng_act, key_act = jax.random.split(agents[i].rng)
                agents[i] = agents[i]._replace(rng=rng_act)
                action = dist.sample(seed=key_act)
                action = np.array(action).squeeze(0)
                if global_step < cfg.START_STEPS: action = env.action_space.sample()
                actions.append(action)
            
            actions = np.array(actions)
            step_res = env.step(actions)
            obs_next, r, done, _, _ = unwrap_step(step_res)
            
            obs_next = np.array(obs_next, dtype=np.float32)
            r = np.array(r, dtype=np.float32).flatten()
            if isinstance(done, (bool, int, np.bool_, np.integer)): d = np.array([float(done)]*num_agents)
            else: d = np.array(done, dtype=np.float32)

            buffer.add(obs, actions, r, d, obs_next)
            obs = obs_next
            ep_ret += r
            
            if buffer.size >= cfg.BATCH_SIZE and global_step > cfg.START_STEPS:
                batch = buffer.sample(cfg.BATCH_SIZE)
                agents = update_step(agents, batch, cfg_dict, global_step)

            if np.any(d): break
        
        # 记录本回合奖励
        total_ep_reward = np.sum(ep_ret)
        total_rewards_history.append(total_ep_reward)
        leader_rewards_history.append(ep_ret[0])
        
        if num_agents > 1:
            # 如果有多个跟随者，取平均值；如果只有一个，直接取值
            follower_ret = np.mean(ep_ret[1:])
            follower_rewards_history.append(follower_ret)
        else:
            follower_rewards_history.append(np.nan)
        
        if (ep + 1) % 10 == 0:
            print(f"Episode {ep+1:04d} | Reward: {total_ep_reward:.2f}")
        
        if (ep + 1) % cfg.SAVE_EVERY == 0:
            save_path = os.path.join(cfg.SAVE_DIR, f"dsact_jax_ep{ep+1}.pkl")
            save_model(agents, save_path)

    final_path = os.path.join(cfg.SAVE_DIR, "dsact_jax_final.pkl")
    save_model(agents, final_path)


# ---- [新增] 导出每个 Episode 的奖励到 Excel ----
    if getattr(cfg, "EXPORT_REWARDS_EXCEL", True):
        excel_name = getattr(cfg, "REWARDS_EXCEL_NAME", "rewards_per_episode.xlsx")
        excel_path = os.path.join(cfg.SAVE_DIR, excel_name)
        try:
            export_rewards_to_excel(
                excel_path,
                total_rewards_history,
                leader_rewards_history,
                follower_rewards_history,
            )
        except Exception as e:
            print(f"[WARN] 导出 Excel 失败: {e}")

    
    # ---- [新增] 绘制平滑奖励曲线 ----
        print("\n" + "="*50)
        print("开始绘制奖励曲线...")
        try:
        # 准备数据
            leader_arr = np.array(leader_rewards_history, dtype=np.float32)
            total_arr = np.array(total_rewards_history, dtype=np.float32)
            follower_arr = None if (num_agents <= 1) else np.array(follower_rewards_history, dtype=np.float32)
        
            method = (getattr(cfg, 'SMOOTH_METHOD', 'gaussian') or '').lower()
            passes = getattr(cfg, 'SMOOTH_PASSES', 1)
            sigma = getattr(cfg, 'SMOOTH_SIGMA', 5.0)
            alpha = getattr(cfg, 'SMOOTH_ALPHA', 0.1)
            window = getattr(cfg, 'SMOOTH_WINDOW', 10)

        # 平滑处理
            if method and method != 'none':
                total_smooth = smooth_curve(total_arr, method, sigma, alpha, window, passes)
                leader_smooth = smooth_curve(leader_arr, method, sigma, alpha, window, passes)
                if follower_arr is not None:
                    follower_smooth = smooth_curve(follower_arr, method, sigma, alpha, window, passes)
            else:
                total_smooth, leader_smooth = total_arr, leader_arr
                follower_smooth = follower_arr

        # 1. 绘制总图 (reward_curves.png)
            fig_path = os.path.join(cfg.SAVE_DIR, "reward_curves.png")
            plt.figure(figsize=(10, 5), dpi=150)
        # 平滑线
            plt.plot(total_smooth, label=f"Total(smooth:{method})", linewidth=2.0, color='green')
            plt.plot(leader_smooth, label=f"Leader(smooth:{method})", linewidth=2.0, color='blue')
            if follower_smooth is not None:
                plt.plot(follower_smooth, label=f"Follower(smooth:{method})", linewidth=2.0, color='red')
        # 原始线 (背景)
            plt.plot(total_arr, linewidth=1.0, color='green', alpha=0.25)
            plt.plot(leader_arr, linewidth=1.0, color='blue', alpha=0.25)
            if follower_arr is not None:
                plt.plot(follower_arr, linewidth=1.0, color='red', alpha=0.25)
        
            plt.xlabel("Episode")
            plt.ylabel("Reward")
            plt.title(f"Training Reward Curves (smoothed: {method})")
            plt.legend()
            plt.grid(True, alpha=0.3)
            plt.tight_layout()
            plt.savefig(fig_path, bbox_inches="tight", dpi=150)
            plt.close()
            print(f"✓ 综合奖励曲线已保存至: {fig_path}")

        # 2. 绘制单独曲线 (Total, Leader, Follower)
            def save_single_curve(data_smooth, data_raw, label, color, filename):
                plt.figure(figsize=(10, 5), dpi=150)
                plt.plot(data_smooth, label=f"{label} (smooth)", linewidth=2.0, color=color)
                if data_raw is not None:
                    plt.plot(data_raw, linewidth=1.0, color=color, alpha=0.25)
                plt.xlabel("Episode")
                plt.ylabel("Reward")
                plt.title(f"{label} Reward Curve")
                plt.legend()
                plt.grid(True, alpha=0.3)
                plt.tight_layout()
                path = os.path.join(cfg.SAVE_DIR, filename)
                plt.savefig(path, bbox_inches="tight", dpi=150)
                plt.close()
                print(f"✓ {label} 曲线已保存至: {path}")

            save_single_curve(total_smooth, total_arr, "Total", "green", "reward_total_only.png")
            save_single_curve(leader_smooth, leader_arr, "Leader", "blue", "reward_leader_only.png")
            if follower_smooth is not None:
                save_single_curve(follower_smooth, follower_arr, "Follower", "red", "reward_follower_only.png")

        except Exception as e:
            print(f"[ERROR] 绘图失败: {e}")
            import traceback
            traceback.print_exc()

        print("="*50 + "\n")
    
        env.close()

# def test(cfg: Config):
#     print(f"进入测试模式: Loading from {cfg.TEST_MODEL_PATH}")
#     if not os.path.exists(cfg.TEST_MODEL_PATH):
#         print(f"Error: 模型文件不存在 {cfg.TEST_MODEL_PATH}")
#         return

#     env = RlGame(n=cfg.N_AGENT, m=cfg.M_ENEMY, render=cfg.TEST_RENDER)
#     if hasattr(env, 'unwrapped'): env = env.unwrapped
    
#     obs = env.reset()
#     obs = np.array(obs, dtype=np.float32)
#     num_agents = cfg.N_AGENT + cfg.M_ENEMY
#     obs_dim = obs.shape[-1]
#     act_dim = env.action_space.shape[0]

#     # 初始化空网络
#     rng = jax.random.PRNGKey(cfg.SEED)
#     agents = []
#     for i in range(num_agents):
#         rng, key = jax.random.split(rng)
#         agents.append(create_agent_state(key, obs_dim, act_dim, cfg, num_agents))
    
#     # 加载权重
#     agents = load_model(agents, cfg.TEST_MODEL_PATH)
    
#     # 性能指标统计
#     total_episodes = cfg.TEST_EPISODES
#     success_count = 0
#     total_fkr = 0.0          # Formation Keeping Rate
#     total_time = 0.0         # Shortest Time (sum of successful steps)
#     total_dist = 0.0         # Shortest Path (sum of trajectory length)
#     total_energy = 0.0       # Minimum Energy (sum of control input)
    
#     success_episodes_count = 0 # 记录成功的总次数用于计算特定指标平均值

#     print(f"\n开始测试 {total_episodes} 回合...")
#     for i in range(total_episodes):
#         obs = env.reset()
#         obs = np.array(obs, dtype=np.float32)
        
#         ep_steps = 0
#         ep_dist = 0.0
#         ep_energy = 0.0
        
#         last_pos = [ob[:2] * 1000.0 for ob in obs] 
#         episode_team_counter = 0

#         for t in range(cfg.EP_LEN):
#             actions = []
#             for j in range(num_agents):
#                 obs_tensor = jnp.expand_dims(jnp.array(obs[j]), 0)
#                 dist = agents[j].actor.apply_fn(agents[j].actor.params, obs_tensor)
#                 rng, key = jax.random.split(rng)
#                 action = dist.sample(seed=key) 
#                 actions.append(np.array(action).squeeze(0))
            
#             actions_np = np.array(actions)
#             step_energy = np.sum(np.abs(actions_np))
#             ep_energy += step_energy

#             step_res = env.step(actions_np)
#             obs_next, r, done, win, team_ctr = unwrap_step(step_res)
#             episode_team_counter = team_ctr 

#             obs_next = np.array(obs_next, dtype=np.float32)
#             curr_pos = [ob[:2] * 1000.0 for ob in obs_next]
            
#             step_dist = 0
#             for k in range(num_agents):
#                 d = np.linalg.norm(curr_pos[k] - last_pos[k])
#                 step_dist += d
#             ep_dist += step_dist
            
#             obs = obs_next
#             last_pos = curr_pos
#             ep_steps += 1
            
#             if cfg.TEST_RENDER:
#                 env.render()
            
#             if np.any(done):
#                 if win:
#                     success_count += 1
#                     success_episodes_count += 1
#                     total_time += ep_steps
#                     total_dist += ep_dist
#                     total_energy += ep_energy
#                 break
        
#         fkr = episode_team_counter / max(ep_steps, 1)
#         total_fkr += fkr

#         print(f"Test Ep {i+1} | Steps: {ep_steps} | Dist: {ep_dist:.1f} | Energy: {ep_energy:.1f} | Win: {win} | FKR: {fkr:.2%}")

#     env.close()

#     # --- 计算统计指标 ---
#     mcr = (success_count / total_episodes) * 100
#     avg_fkr = (total_fkr / total_episodes) * 100
    
#     if success_episodes_count > 0:
#         avg_time = total_time / success_episodes_count
#         avg_dist = total_dist / success_episodes_count
#         avg_energy = total_energy / success_episodes_count
#     else:
#         avg_time = 0
#         avg_dist = 0
#         avg_energy = 0

#     print("\n" + "="*40)
#     print("           测试结果汇总            ")
#     print("="*40)
#     print(f"任务完成率 (MCR):            {mcr:.2f}%")
#     print(f"平均编队保持率 (FKR):        {avg_fkr:.2f}%")
#     print(f"平均飞行时间 (Steps):        {avg_time:.2f}")
#     print(f"平均飞行路程 (Distance):     {avg_dist:.2f}")
#     print(f"平均能量损耗 (Energy):       {avg_energy:.2f}")
#     print("="*40)

def test(cfg: Config):
    print(f"进入测试模式: Loading from {cfg.TEST_MODEL_PATH}")
    if not os.path.exists(cfg.TEST_MODEL_PATH):
        print(f"Error: 模型文件不存在 {cfg.TEST_MODEL_PATH}")
        return

    env = RlGame(n=cfg.N_AGENT, m=cfg.M_ENEMY, render=cfg.TEST_RENDER)
    if hasattr(env, 'unwrapped'): env = env.unwrapped
    
    obs = env.reset()
    obs = np.array(obs, dtype=np.float32)
    num_agents = cfg.N_AGENT + cfg.M_ENEMY
    obs_dim = obs.shape[-1]
    act_dim = env.action_space.shape[0]

    # 初始化空网络
    rng = jax.random.PRNGKey(cfg.SEED)
    agents = []
    for i in range(num_agents):
        rng, key = jax.random.split(rng)
        agents.append(create_agent_state(key, obs_dim, act_dim, cfg, num_agents))
    
    # 加载权重
    agents = load_model(agents, cfg.TEST_MODEL_PATH)
    
    # === [关键修改] 对齐 main_SAC_dsact_test.py 的统计变量 ===
    win_times = 0
    total_fkr = 0.0          
    total_steps = 0.0         
    total_integral_V = 0.0   # 飞行路程 (速度积分)
    total_integral_U = 0.0   # 能量损耗 (动作值积分)
    
    total_episodes = cfg.TEST_EPISODES

    print(f"\n开始测试 {total_episodes} 回合 (计算逻辑已对齐 main_SAC_dsact_test.py)...")
    
    for i in range(total_episodes):
        obs = env.reset()
        obs = np.array(obs, dtype=np.float32)
        
        # 单回合累加器
        ep_integral_V = 0.0
        ep_integral_U = 0.0
        last_team_counter = 0.0
        steps_taken = 0

        for t in range(cfg.EP_LEN):
            actions = []
            for j in range(num_agents):
                obs_tensor = jnp.expand_dims(jnp.array(obs[j]), 0)
                dist = agents[j].actor.apply_fn(agents[j].actor.params, obs_tensor)
                rng, key = jax.random.split(rng)
                # 测试时通常使用确定性策略 (Mean)，但这里保持采样以防分布偏移
                # 为了完全对齐测试脚本的 deterministic 逻辑，其实应该取 mean，
                # 但 JAX TanhDist 取 mean 比较麻烦，sample 方差很小，近似均值。
                action = dist.sample(seed=key) 
                actions.append(np.array(action).squeeze(0))
            
            actions_np = np.array(actions)
            
            step_res = env.step(actions_np)
            obs_next, r, done, win, team_ctr = unwrap_step(step_res)
            
            obs_next = np.array(obs_next, dtype=np.float32)
            
            # === [关键修改] 模仿测试脚本的累加逻辑 ===
            # 1. 飞行路程 = 领机速度积分 (obs[0][2] 是领机速度)
            ep_integral_V += float(obs[0][2])
            
            # 2. 能量损耗 = 领机动作绝对值之和 (只看领机)
            ep_integral_U += float(np.abs(actions_np[0]).sum())
            
            obs = obs_next
            last_team_counter = float(team_ctr)
            steps_taken = t + 1
            
            if cfg.TEST_RENDER:
                env.render()
            
            if np.any(done):
                if win:
                    win_times += 1
                break
        
        # 编队保持率计算
        fkr = last_team_counter / max(steps_taken, 1)
        
        # 累加到总统计中
        total_fkr += fkr
        total_steps += steps_taken
        total_integral_V += ep_integral_V
        total_integral_U += ep_integral_U

        print(f"Test Ep {i+1:03d} | Steps: {steps_taken:4d} | Dist(V): {ep_integral_V:.2f} | Energy(U): {ep_integral_U:.2f} | Win: {bool(win)} | FKR: {fkr:.4f}")

    env.close()

    # --- [关键修改] 最终计算方式对齐 (全部除以总回合数) ---
    avg_mcr = win_times / total_episodes
    avg_fkr = total_fkr / total_episodes
    avg_time = total_steps / total_episodes
    avg_dist = total_integral_V / total_episodes
    avg_energy = total_integral_U / total_episodes

    print("\n" + "="*40)
    print("           测试结果汇总 (Aligned)            ")
    print("="*40)
    print(f"任务完成率 (MCR):            {avg_mcr:.4f}")
    print(f"平均最大编队保持率 (FKR):    {avg_fkr:.4f}")
    print(f"平均最短飞行时间 (Steps):    {avg_time:.4f}")
    print(f"平均最短飞行路程 (Dist):     {avg_dist:.4f}")
    print(f"平均最小能量损耗 (Energy):   {avg_energy:.4f}")
    print("="*40)

if __name__ == "__main__":
    os.environ["XLA_PYTHON_CLIENT_PREALLOCATE"] = "false"
    cfg = Config()
    
    if cfg.TEST_MODE:
        test(cfg)
    else:
        train(cfg)